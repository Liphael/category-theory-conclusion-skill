#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 extract_text.py 的四格式测试夹具（仅用于门禁自测，不参与技能运行）。

用法：
    python _make_fixtures.py --out .selftest/fixtures              # 7 个（含 2 个异常夹具）
    python _make_fixtures.py --out .selftest/fixtures4 --skip-bad  # 5 个正常夹具
"""
import os
import argparse

p = argparse.ArgumentParser()
p.add_argument("--out", default="./fixtures", help="夹具输出目录")
p.add_argument("--skip-bad", action="store_true",
               help="跳过异常夹具（legacy.doc 不支持格式 / scanned_no_text.pdf 无文本层）")
args = p.parse_args()

OUT = args.out
os.makedirs(OUT, exist_ok=True)

# 1) TXT · UTF-8
with open(os.path.join(OUT, "note_utf8.txt"), "w", encoding="utf-8") as f:
    f.write("总体与样本\n总体：研究对象的全体。\n样本：总体中的一个子集。\n参数：总体的数值特征。\n统计量：样本的数值特征。\n")

# 2) TXT · GBK（测试编码回退）
with open(os.path.join(OUT, "note_gbk.txt"), "w", encoding="gbk") as f:
    f.write("编码回退测试：这一段用 GBK 编码写入，用于验证提取器的编码探测链路是否生效。\n")

# 3) Markdown（含 frontmatter、表格、代码围栏）
with open(os.path.join(OUT, "lecture.md"), "w", encoding="utf-8") as f:
    f.write("---\ntitle: 假设检验\nweek: 3\n---\n\n# 假设检验\n\n## 基本术语\n\n"
            "- 原假设 H0\n- 备择假设 Ha\n\n| 错误类型 | 含义 |\n| --- | --- |\n"
            "| 第一类 | H0 真却拒绝 |\n| 第二类 | H0 假却未拒绝 |\n\n```r\n"
            "t.test(x, mu = 120)\n```\n")

# 4) DOCX（段落 + 表格）
from docx import Document
d = Document()
d.add_heading("方差分析 ANOVA", level=1)
d.add_paragraph("总变异可分解为组间变异与组内变异两部分。")
d.add_paragraph("Total SS = SST + SSE")
t = d.add_table(rows=3, cols=3)
t.style = "Table Grid"
hdr = ["来源", "平方和", "自由度"]
for i, h in enumerate(hdr):
    t.rows[0].cells[i].text = h
for r, row in enumerate([("组间", "SST", "k-1"), ("组内", "SSE", "N-k")], start=1):
    for c, v in enumerate(row):
        t.rows[r].cells[c].text = v
d.save(os.path.join(OUT, "report.docx"))

# 5) PDF（3 页，含文本）
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
c = canvas.Canvas(os.path.join(OUT, "slides.pdf"), pagesize=A4)
for i in range(1, 4):
    c.setFont("Helvetica", 14)
    c.drawString(72, 760, "Lecture Slide - Page %d" % i)
    c.setFont("Helvetica", 11)
    c.drawString(72, 730, "Central Limit Theorem: Xbar ~ N(mu, sigma^2/n) as n grows.")
    c.drawString(72, 710, "Standardisation: Z = (Xbar - mu) / (sigma / sqrt(n)) ~ N(0,1).")
    c.showPage()
c.save()

if args.skip_bad:
    print("fixtures ->", os.path.abspath(OUT), "(skip-bad)")
    for fn in sorted(os.listdir(OUT)):
        print("  ", fn, os.path.getsize(os.path.join(OUT, fn)), "bytes")
    raise SystemExit(0)

# 6) PDF（无文本层，模拟扫描件）
c2 = canvas.Canvas(os.path.join(OUT, "scanned_no_text.pdf"), pagesize=A4)
c2.setFillColorRGB(0.85, 0.85, 0.85)
c2.rect(72, 400, 400, 300, fill=1, stroke=0)
c2.showPage()
c2.save()

# 7) 旧版 .doc（不支持格式，测试告警）
with open(os.path.join(OUT, "legacy.doc"), "wb") as f:
    f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 200)

print("fixtures ->", os.path.abspath(OUT))
for fn in sorted(os.listdir(OUT)):
    print("  ", fn, os.path.getsize(os.path.join(OUT, fn)), "bytes")
