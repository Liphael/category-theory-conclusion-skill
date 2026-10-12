#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_text.py —— 材料统一文本提取器（范畴论总结技能 · 步骤 0）

支持输入格式：.pdf / .docx / .txt / .md（及 .markdown）
逐文件按扩展名路由，输出纯文本到 --out 目录，并生成 manifest.json。

设计红线（不可违反）：
  1. 不静默丢弃：提取为空 / 格式不支持 / 读取失败，都必须写进 manifest 并在摘要中列出。
  2. 不猜测内容：只做格式转换，不改写、不摘要、不补全。
  3. 路径与编码自决：编码按 utf-8-sig → utf-8 → gb18030 → latin-1 顺序探测，不询问用户。

用法：
  python extract_text.py --input <材料目录> --out <输出目录> [--min-chars 50] [--recursive]

退出码：
  0 = 全部文件提取成功（虽有文件低于 --min-chars 会告警但仍是 0）
  1 = 存在 empty / unsupported / error 状态的文件（门禁应判失败）
  2 = 参数错误（输入目录不存在等）
"""

import argparse
import json
import os
import sys
import traceback

TEXT_EXTS = {".txt", ".md", ".markdown"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
KNOWN_BUT_UNSUPPORTED = {".doc": "旧版 Word 二进制格式，需先另存为 .docx",
                         ".rtf": "富文本格式，需先另存为 .docx 或 .txt",
                         ".pages": "Apple Pages 格式，需先导出为 .docx 或 .txt",
                         ".html": "网页格式，需先转存为 .md 或 .txt",
                         ".htm": "网页格式，需先转存为 .md 或 .txt"}
ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "latin-1")
SUPPORTED = sorted(TEXT_EXTS | PDF_EXTS | DOCX_EXTS)


def read_text_file(path):
    """按编码顺序探测读取纯文本/markdown。返回 (文本, 使用的编码) 或抛异常。"""
    raw = open(path, "rb").read()
    if not raw.strip():
        return "", ""
    last_err = None
    for enc in ENCODINGS:
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError as e:
            last_err = e
    raise last_err


def extract_pdf(path):
    """PDF → 文本。逐页标注 PAGE i/N。返回 (文本, 页数, 正文字符数, 警告列表)。

    注意：正文字符数 body_chars 只统计页内文本，不含注入的 PAGE 页眉标记，
    否则空页 PDF 会因页眉而误判为「有内容」（假阴性）。
    """
    from pypdf import PdfReader
    reader = PdfReader(path)
    warnings = []
    chunks = []
    total = len(reader.pages)
    empty_pages = 0
    body_chars = 0
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception as e:  # 单页失败不整体中止
            t = ""
            warnings.append("第 %d 页提取失败：%s" % (i, type(e).__name__))
        body_chars += len(t.strip())
        if not t.strip():
            empty_pages += 1
        chunks.append("\n===== [%s] PAGE %d/%d =====\n%s\n"
                      % (os.path.basename(path), i, total, t))
    if empty_pages:
        warnings.append("%d/%d 页无文本层（疑似扫描件或图片页），需 OCR 或人工确认"
                        % (empty_pages, total))
    return "".join(chunks), total, body_chars, warnings


def extract_docx(path):
    """DOCX → 文本。段落顺序保留，表格以 [TABLE n] 标注。
    返回 (文本, 段落数, 表格数, 正文字符数, 警告)。"""
    import docx
    warnings = []
    d = docx.Document(path)
    out = ["\n===== [%s] DOCX BODY =====\n" % os.path.basename(path)]
    para_n = 0
    body_chars = 0
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            out.append(t + "\n")
            para_n += 1
            body_chars += len(t)
    table_n = 0
    for ti, tb in enumerate(d.tables, 1):
        table_n += 1
        out.append("\n----- [%s] TABLE %d -----\n" % (os.path.basename(path), ti))
        for row in tb.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            line = " | ".join(cells)
            out.append(line + "\n")
            body_chars += len(line)
        out.append("----- END TABLE %d -----\n" % ti)
    # 页眉/页脚不在正文流中：仅在确有非空文本时告警，避免空段落造成假阳性
    try:
        hf = []
        for s in d.sections:
            for part in (s.header, s.footer):
                for p in part.paragraphs:
                    if p.text.strip():
                        hf.append(p.text.strip())
        if hf:
            warnings.append("页眉/页脚含 %d 段文本，python-docx 默认不提取；如影响结论请说明"
                            % len(hf))
    except Exception:
        pass
    return "".join(out), para_n, table_n, body_chars, warnings


def extract_plain(path):
    """TXT / MD → 文本。返回 (文本, 编码, 正文字符数, 警告)。"""
    warnings = []
    text, enc = read_text_file(path)
    ext = os.path.splitext(path)[1].lower()
    out = "\n===== [%s] %s =====\n%s\n" % (
        os.path.basename(path), "MARKDOWN" if ext != ".txt" else "TEXT", text)
    body_chars = len(text.strip())
    if body_chars == 0:
        warnings.append("文件为空或仅含空白字符")
    else:
        warnings.append("读取编码：%s" % enc)
    return out, enc, body_chars, warnings


def iter_sources(input_dir, recursive):
    for root, dirs, files in os.walk(input_dir):
        if not recursive and root != input_dir:
            continue
        dirs.sort()
        for f in sorted(files):
            if f.startswith("~$") or f.startswith("."):
                continue
            yield os.path.join(root, f)


def main():
    ap = argparse.ArgumentParser(description="材料统一文本提取器")
    ap.add_argument("--input", required=True, help="材料目录")
    ap.add_argument("--out", required=True, help="文本输出目录")
    ap.add_argument("--min-chars", type=int, default=50,
                    help="低于此字符数视为可疑（默认 50），仅告警")
    ap.add_argument("--recursive", action="store_true", help="递归子目录")
    ap.add_argument("--manifest", default=None,
                    help="manifest.json 路径（默认 <out>/manifest.json）")
    args = ap.parse_args()

    if not os.path.isdir(args.input):
        print("[FATAL] 输入目录不存在：%s" % args.input, file=sys.stderr)
        return 2
    os.makedirs(args.out, exist_ok=True)

    manifest_path = args.manifest or os.path.join(args.out, "manifest.json")
    records = []

    for path in iter_sources(args.input, args.recursive):
        ext = os.path.splitext(path)[1].lower()
        rel = os.path.relpath(path, args.input)
        rec = {"path": rel, "ext": ext, "status": "ok", "chars": 0,
               "pages": None, "paragraphs": None, "tables": None,
               "encoding": None, "output": None, "warnings": []}
        try:
            body_chars = 0
            if ext in PDF_EXTS:
                text, n, body_chars, warns = extract_pdf(path)
                rec["pages"] = n
                rec["warnings"] += warns
            elif ext in DOCX_EXTS:
                text, pn, tn, body_chars, warns = extract_docx(path)
                rec["paragraphs"] = pn
                rec["tables"] = tn
                rec["warnings"] += warns
            elif ext in TEXT_EXTS:
                text, enc, body_chars, warns = extract_plain(path)
                rec["encoding"] = enc
                rec["warnings"] += warns
            elif ext in KNOWN_BUT_UNSUPPORTED:
                rec["status"] = "unsupported"
                rec["warnings"].append(KNOWN_BUT_UNSUPPORTED[ext])
                records.append(rec)
                continue
            else:
                rec["status"] = "skipped"
                rec["warnings"].append("非材料扩展名，跳过")
                records.append(rec)
                continue

            text = text or ""
            rec["chars"] = body_chars  # 只计正文，排除注入的页眉标记
            if rec["chars"] == 0:
                rec["status"] = "empty"
                rec["warnings"].append("正文提取为空，无法进入总结流程")
            elif rec["chars"] < args.min_chars:
                rec["warnings"].append("正文字符数 %d 低于阈值 %d，内容可能不完整"
                                       % (rec["chars"], args.min_chars))

            out_name = rel.replace(os.sep, "__")
            out_name = os.path.splitext(out_name)[0] + ".txt"
            out_path = os.path.join(args.out, out_name)
            with open(out_path, "w", encoding="utf-8") as fh:
                fh.write(text)
            rec["output"] = out_path
        except Exception as e:
            rec["status"] = "error"
            rec["warnings"].append("%s: %s" % (type(e).__name__, e))
            rec["traceback"] = traceback.format_exc(limit=2)
        records.append(rec)

    with open(manifest_path, "w", encoding="utf-8") as fh:
        json.dump({"input_dir": os.path.abspath(args.input),
                   "out_dir": os.path.abspath(args.out),
                   "min_chars": args.min_chars,
                   "supported_exts": SUPPORTED,
                   "files": records}, fh, ensure_ascii=False, indent=2)

    # ---- 摘要表（硬证据，可复跑） ----
    print("=" * 78)
    print("提取摘要　输入=%s" % os.path.abspath(args.input))
    print("=" * 78)
    print("%-34s %-6s %-9s %8s  %s" % ("文件", "格式", "状态", "字符数", "备注"))
    print("-" * 78)
    bad = []
    for r in records:
        note = "；".join(r["warnings"])[:44] if r["warnings"] else ""
        extra = ""
        if r["pages"] is not None:
            extra = "%d页 " % r["pages"]
        if r["tables"]:
            extra += "%d表 " % r["tables"]
        print("%-34s %-6s %-9s %8d  %s%s" % (r["path"][:34], r["ext"], r["status"],
                                             r["chars"], extra, note))
        if r["status"] in ("empty", "unsupported", "error"):
            bad.append(r)
    print("-" * 78)
    ok = [r for r in records if r["status"] == "ok"]
    print("合计 %d 个文件：成功 %d，可疑/失败 %d" % (len(records), len(ok), len(bad)))
    print("manifest: %s" % manifest_path)
    if bad:
        print("\n[需要处置] 以下文件未成功提取，必须先处置再进入步骤 1：")
        for r in bad:
            print("  - %s (%s) ：%s" % (r["path"], r["status"], "；".join(r["warnings"])))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
